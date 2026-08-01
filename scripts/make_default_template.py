"""Generate the shipped default docxtpl template (src/docgen/templates/default.docx).

The template is the docx shell every generated document is poured into. A
branded client template must provide the same two jinja placeholders:

    {{ title }}     — document title (put on the cover / first page)
    {{ subtitle }}  — solution name, version and generation note

The generated body is appended after the template content using the
template's styles (Heading 1-4, List Bullet, Table Grid).

Run:  python scripts/make_default_template.py
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

TARGET = Path(__file__).parent.parent / "src" / "docgen" / "templates" / "default.docx"


def main() -> None:
    doc = Document()

    title = doc.add_paragraph("{{ title }}", style="Title")
    title.runs[0].font.size = Pt(32)

    subtitle = doc.add_paragraph("{{ subtitle }}", style="Subtitle")
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_paragraph()  # spacer before the generated body

    # Make sure the table style the emitter uses is embedded in the template.
    seed = doc.add_table(rows=1, cols=1)
    seed.style = "Table Grid"
    seed._element.getparent().remove(seed._element)

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TARGET))
    print(f"wrote {TARGET}")


if __name__ == "__main__":
    main()
