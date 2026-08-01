"""DocModel → docx emitter via docxtpl.

The template (shipped default or a branded replacement configured with
`docx_template`) provides the shell: `{{ title }}` / `{{ subtitle }}`
placeholders, styles, headers/footers. After docxtpl renders those
placeholders, the generated body is appended to the end of the template
document using its styles — so a branded template only needs the two
placeholders plus (optionally) restyled Heading/Table Grid styles.
See docs/USAGE.md.
"""

from __future__ import annotations

from importlib import resources
from pathlib import Path

from docx.shared import Pt, RGBColor
from docxtpl import DocxTemplate

from docgen.renderers.diagrams import DiagramService
from docgen.renderers.docmodel import (
    BulletList,
    Callout,
    CodeBlock,
    Diagram,
    Document,
    Paragraph,
    Placeholder,
    Section,
    Table,
)

_CALLOUT_COLORS = {"info": RGBColor(0x1F, 0x4E, 0x79), "warning": RGBColor(0x9C, 0x57, 0x00), "risk": RGBColor(0xC0, 0x00, 0x00)}
_CALLOUT_LABELS = {"info": "Note", "warning": "Warning", "risk": "Risk"}


def default_template_path() -> Path:
    with resources.as_file(resources.files("docgen") / "templates" / "default.docx") as p:
        return Path(p)


def _add_code_block(container, text: str) -> None:
    for line in text.split("\n"):
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8)


def _heading_style(level: int) -> str:
    return f"Heading {min(level, 4)}"


def _emit_block(block, container, level: int, diagrams: DiagramService) -> None:
    if isinstance(block, Section):
        container.add_paragraph(block.heading, style=_heading_style(level))
        for child in block.blocks:
            _emit_block(child, container, level + 1, diagrams)
    elif isinstance(block, Paragraph):
        p = container.add_paragraph()
        run = p.add_run(block.text)
        run.bold = block.bold
    elif isinstance(block, BulletList):
        for item in block.items:
            container.add_paragraph(item, style="List Bullet")
    elif isinstance(block, Table):
        if block.caption:
            cap = container.add_paragraph()
            cap.add_run(block.caption).italic = True
        table = container.add_table(rows=1, cols=len(block.headers))
        table.style = "Table Grid"
        header_cells = table.rows[0].cells
        for i, header in enumerate(block.headers):
            header_cells[i].text = ""
            run = header_cells[i].paragraphs[0].add_run(str(header))
            run.bold = True
        for row in block.rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                if i < len(cells):
                    cells[i].text = str(value)
        container.add_paragraph()
    elif isinstance(block, CodeBlock):
        _add_code_block(container, block.text)
        container.add_paragraph()
    elif isinstance(block, Diagram):
        if block.caption:
            cap = container.add_paragraph()
            cap.add_run(block.caption).italic = True
        image = diagrams.render_image(block)
        if image is not None:
            container.add_picture(str(image))
        else:
            _add_code_block(container, block.mermaid)
            note = container.add_paragraph()
            note.add_run(diagrams.fallback_note or "Diagram source (no diagram renderer available).").italic = True
        container.add_paragraph()
    elif isinstance(block, Placeholder):
        p = container.add_paragraph()
        run = p.add_run(block.text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        if block.hint:
            p.add_run(f" — {block.hint}").italic = True
    elif isinstance(block, Callout):
        p = container.add_paragraph()
        label_run = p.add_run(f"{_CALLOUT_LABELS.get(block.severity, block.severity.title())}: ")
        label_run.bold = True
        color = _CALLOUT_COLORS.get(block.severity)
        if color is not None:
            label_run.font.color.rgb = color
        p.add_run(block.text)
    else:  # pragma: no cover - defensive
        raise TypeError(f"Unknown block type: {type(block).__name__}")


def write_docx(document: Document, path: Path, diagrams: DiagramService,
               template_path: Path | None = None) -> Path:
    template = Path(template_path) if template_path else default_template_path()
    tpl = DocxTemplate(str(template))
    tpl.render({"title": document.title, "subtitle": document.subtitle})
    body = tpl.docx  # the underlying python-docx Document; body is appended after template content
    for section in document.sections:
        _emit_block(section, body, 1, diagrams)
    path.parent.mkdir(parents=True, exist_ok=True)
    tpl.save(str(path))
    return path
