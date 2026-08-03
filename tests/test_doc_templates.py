"""Per-document Word templates: HLD.docx drives the HLD, LLD.docx the LLD, and
a missing template silently falls back to the existing behaviour."""

import shutil
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from docgen.commands import render_documents, render_transcript_documents
from docgen.config import DocgenConfig, load_config
from docgen.doc_templates import candidate_names, find_document_template, resolve_template
from docgen.parsers import parse_solution
from docgen.renderers.docx import default_template_path
from docgen.rules_io import load_rules
from docgen.transcripts import parse_transcript

TRANSCRIPTS = Path(__file__).parent / "fixtures" / "transcripts"


def make_template(directory: Path, name: str, marker: str) -> Path:
    """A real .docx template carrying a recognisable paragraph."""
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / name
    shutil.copy(default_template_path(), path)
    document = DocxDocument(str(path))
    document.paragraphs[0].insert_paragraph_before(marker)
    document.save(str(path))
    return path


@pytest.fixture(scope="module")
def rich(built_fixtures):
    return parse_solution(built_fixtures / "rich.zip")


# ---------------------------------------------------------------------------
# Name matching
# ---------------------------------------------------------------------------


def test_candidate_names_include_key_alias_and_title():
    # "HLD" and the key "hld" normalise to the same thing, so only one is kept
    names = candidate_names("hld", "High-Level Design")
    assert "hld" in names and "High-Level Design" in names
    assert candidate_names("data-dictionary", "Data Dictionary") == ["data-dictionary", "DD"]


@pytest.mark.parametrize("filename", ["HLD.docx", "hld.docx", "High-Level Design.docx",
                                      "high_level_design.docx", "HIGH LEVEL DESIGN.docx"])
def test_template_matched_case_and_separator_insensitively(tmp_path: Path, filename):
    templates = tmp_path / "templates"
    make_template(templates, filename, "BRANDED")
    cfg = DocgenConfig(templates_dir=templates)
    assert find_document_template("hld", "High-Level Design", cfg) is not None


def test_template_for_one_document_does_not_leak_to_another(tmp_path: Path):
    templates = tmp_path / "templates"
    make_template(templates, "HLD.docx", "BRANDED HLD")
    cfg = DocgenConfig(templates_dir=templates)
    assert find_document_template("hld", "High-Level Design", cfg) is not None
    assert find_document_template("lld", "Low-Level Design", cfg) is None


def test_missing_templates_folder_falls_back_to_shipped_default(tmp_path: Path):
    cfg = DocgenConfig(templates_dir=tmp_path / "nope")
    template, note = resolve_template("hld", "High-Level Design", cfg)
    assert template is None and note == ""


def test_per_document_template_beats_the_global_one(tmp_path: Path):
    templates = tmp_path / "templates"
    specific = make_template(templates, "HLD.docx", "BRANDED HLD")
    global_template = make_template(tmp_path, "house-style.docx", "HOUSE STYLE")
    cfg = DocgenConfig(templates_dir=templates, docx_template=global_template)
    assert resolve_template("hld", "High-Level Design", cfg)[0] == specific
    assert resolve_template("lld", "Low-Level Design", cfg)[0] == global_template


# ---------------------------------------------------------------------------
# End to end
# ---------------------------------------------------------------------------


def test_solution_document_rendered_into_its_own_template(rich, tmp_path: Path):
    templates = tmp_path / "templates"
    make_template(templates, "HLD.docx", "HLD HOUSE TEMPLATE")
    make_template(templates, "LLD.docx", "LLD HOUSE TEMPLATE")
    out_dir = tmp_path / "out"
    cfg = DocgenConfig(templates_dir=templates)

    render_documents(rich, ["hld", "lld", "rraid"], ["docx"], out_dir, cfg, no_llm=True)

    hld = "\n".join(p.text for p in DocxDocument(str(out_dir / "hld.docx")).paragraphs)
    lld = "\n".join(p.text for p in DocxDocument(str(out_dir / "lld.docx")).paragraphs)
    rraid = "\n".join(p.text for p in DocxDocument(str(out_dir / "rraid.docx")).paragraphs)
    assert "HLD HOUSE TEMPLATE" in hld and "LLD HOUSE TEMPLATE" not in hld
    assert "LLD HOUSE TEMPLATE" in lld
    assert "HOUSE TEMPLATE" not in rraid  # no template for it — shipped default
    assert "High-Level Design" in hld  # the document itself still renders


def test_transcript_document_rendered_into_its_own_template(tmp_path: Path):
    templates = tmp_path / "templates"
    make_template(templates, "Meeting Notes.docx", "MINUTES HOUSE TEMPLATE")
    out_dir = tmp_path / "out"
    cfg = DocgenConfig(templates_dir=templates)
    cues = load_rules(DocgenConfig()).get("transcript_cues", {})
    snapshot = parse_transcript(TRANSCRIPTS / "workshop.txt", cues)

    render_transcript_documents(snapshot, ["meeting-notes", "actions"], ["docx"], out_dir,
                                cfg, no_llm=True)

    notes = "\n".join(p.text for p in DocxDocument(str(out_dir / "meeting-notes.docx")).paragraphs)
    actions = "\n".join(p.text for p in DocxDocument(str(out_dir / "actions.docx")).paragraphs)
    assert "MINUTES HOUSE TEMPLATE" in notes
    assert "MINUTES HOUSE TEMPLATE" not in actions


def test_templates_dir_configurable_and_defaults_to_templates(tmp_path: Path):
    assert DocgenConfig().templates_dir == Path("templates")
    config_file = tmp_path / "docgen.yaml"
    config_file.write_text("templates_dir: brand/word\n", encoding="utf-8")
    assert load_config(config_file).templates_dir == Path("brand/word")
