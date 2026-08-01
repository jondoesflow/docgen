"""Renderer tests: every registered document renders from every fixture in both
formats, offline (--no-llm). Extended automatically as renderers register."""

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from docgen.commands import render_documents
from docgen.config import DocgenConfig
from docgen.parsers import parse_solution
from docgen.renderers import DOC_RENDERERS

FIXTURE_ZIPS = ["minimal", "rich", "diff_v1", "diff_v2"]


@pytest.fixture(scope="module")
def snapshots(built_fixtures):
    return {name: parse_solution(built_fixtures / f"{name}.zip") for name in FIXTURE_ZIPS}


@pytest.mark.parametrize("fixture_name", FIXTURE_ZIPS)
def test_all_docs_render_both_formats(fixture_name, snapshots, tmp_path: Path):
    snapshot = snapshots[fixture_name]
    keys = list(DOC_RENDERERS)
    written = render_documents(snapshot, keys, ["md", "docx"], tmp_path, DocgenConfig(), no_llm=True)
    applicable = [k for k in keys if DOC_RENDERERS[k]().applies(snapshot)]
    assert len(written) == 2 * len(applicable)
    for path in written:
        assert path.is_file()
        assert path.stat().st_size > 200, f"{path.name} suspiciously small"
        if path.suffix == ".docx":
            docx = DocxDocument(str(path))
            assert len(docx.paragraphs) > 3


def test_lld_markdown_content(snapshots, tmp_path: Path):
    render_documents(snapshots["rich"], ["lld"], ["md"], tmp_path, DocgenConfig(), no_llm=True)
    text = (tmp_path / "lld.md").read_text(encoding="utf-8")
    assert "# Low-Level Design — Rich CE Solution" in text
    assert "abc_project" in text and "abc_task" in text
    assert "```mermaid" in text and "erDiagram" in text
    # relationship drawn in the project cluster ERD
    assert 'abc_project ||--o{ abc_task : "abc_project_task"' in text
    # two functional clusters: project group and lone invoice
    assert text.count("Functional cluster:") == 2
    # component inventory counts include not-yet-parsed families as other components
    assert "Component inventory" in text


def test_data_dictionary_flags_missing_descriptions(snapshots, tmp_path: Path):
    render_documents(snapshots["rich"], ["data-dictionary"], ["md"], tmp_path, DocgenConfig(), no_llm=True)
    text = (tmp_path / "data-dictionary.md").read_text(encoding="utf-8")
    assert "⚠ No description" in text  # abc_description has no description
    assert "Columns without a description" in text
    assert "Project Status" in text  # global option set section


def test_docx_contains_title_and_tables(snapshots, tmp_path: Path):
    render_documents(snapshots["minimal"], ["lld"], ["docx"], tmp_path, DocgenConfig(), no_llm=True)
    docx = DocxDocument(str(tmp_path / "lld.docx"))
    all_text = "\n".join(p.text for p in docx.paragraphs)
    assert "Low-Level Design — Minimal Solution" in all_text
    assert len(docx.tables) >= 2  # solution facts + at least the widget columns table
    cells = [c.text for t in docx.tables for row in t.rows for c in row.cells]
    assert "min_name" in cells


def test_diagram_fallback_note_when_no_renderer(snapshots, tmp_path: Path, monkeypatch):
    """With neither mmdc nor dot on PATH the docx embeds mermaid source, never fails."""
    import shutil as shutil_mod

    monkeypatch.setattr(shutil_mod, "which", lambda name: None)
    written = render_documents(snapshots["rich"], ["lld"], ["docx"], tmp_path, DocgenConfig(), no_llm=True)
    assert (tmp_path / "lld.docx").is_file()
    assert not (tmp_path / "diagrams").exists()
