"""Transcript renderer matrix: every registered transcript document renders
from every transcript fixture in both formats, offline."""

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from docgen.commands import render_transcript_documents, run_all_transcript
from docgen.config import DocgenConfig
from docgen.renderers import TRANSCRIPT_DOC_RENDERERS
from docgen.renderers.docmodel import PLACEHOLDER_TEXT
from docgen.rules_io import load_rules
from docgen.transcripts import parse_transcript

TRANSCRIPTS = Path(__file__).parent / "fixtures" / "transcripts"
FIXTURES = ["workshop.txt", "teams_export.txt", "teams_meeting.vtt", "plain_notes.txt"]


@pytest.fixture(scope="module")
def snapshots():
    cues = load_rules(DocgenConfig()).get("transcript_cues", {})
    return {name: parse_transcript(TRANSCRIPTS / name, cues) for name in FIXTURES}


@pytest.mark.parametrize("fixture_name", FIXTURES)
def test_all_transcript_docs_render_both_formats(fixture_name, snapshots, tmp_path: Path):
    snapshot = snapshots[fixture_name]
    keys = list(TRANSCRIPT_DOC_RENDERERS)
    written = render_transcript_documents(snapshot, keys, ["md", "docx"], tmp_path,
                                          DocgenConfig(), no_llm=True)
    applicable = [k for k in keys if TRANSCRIPT_DOC_RENDERERS[k]().applies(snapshot)]
    assert len(written) == 2 * len(applicable)
    for path in written:
        assert path.is_file()
        assert path.stat().st_size > 200, f"{path.name} suspiciously small"
        if path.suffix == ".docx":
            docx = DocxDocument(str(path))
            assert len(docx.paragraphs) > 3


def test_meeting_notes_content(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["meeting-notes"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "meeting-notes.md").read_text(encoding="utf-8")
    assert "# Meeting Notes — Discovery Workshop 1 of 2 - Service Operations" in text
    assert "Northgate Utilities plc" in text
    assert "Rowan Ellisdale" in text and "Operations Director" in text
    assert "joining remotely from Hull" in text
    # section structure and the registers
    assert "Job Scheduling" in text and "Customer Notifications" in text
    assert "Send call-reason tagging analysis from autumn 2025" in text
    assert "Sequencing of customer notifications vs engineer mobile app" in text
    # narrative left to the consultant offline
    assert PLACEHOLDER_TEXT in text
    # coverage is stated, never hidden
    assert "Transcript quality and coverage" in text


def test_requirements_catalogue_quotes_and_priorities(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["requirements"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "requirements.md").read_text(encoding="utf-8")
    assert "REQ-001" in text
    assert "Priority read from" in text
    assert "non-negotiable" in text
    # every catalogue row cites a line number
    assert ", line " in text
    assert "Non-functional requirements" in text


def test_actions_register_flags_missing_dates(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["actions"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "actions.md").read_text(encoding="utf-8")
    assert "A3" in text and "⚠ No date" in text
    assert "Actions by owner" in text
    assert "Kaya Petrov" in text


def test_rraid_seeds_and_human_judgement_sections(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["rraid"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "rraid.md").read_text(encoding="utf-8")
    assert "RSK-001" in text
    for heading in ("## Risks", "## Assumptions", "## Issues", "## Dependencies", "## Constraints"):
        assert heading in text
    assert "Additional risks (human judgement)" in text
    # parked items are carried forward as open items with a delivery cost
    assert "Open items carried from the session" in text
    assert "P1" in text


# ---------------------------------------------------------------------------
# Pre-build design documents — same keys as the solution set
# ---------------------------------------------------------------------------


DESIGN_KEYS = ("hld", "lld", "data-dictionary", "security", "deployment",
               "licensing", "integration", "rraid", "hygiene")


def test_design_keys_match_the_solution_document_keys():
    """The point of the design set: same key before build and after."""
    from docgen.constants import ALL_DOC_KEYS, ALL_TRANSCRIPT_DOC_KEYS
    from docgen.renderers import DOC_RENDERERS

    shared = set(ALL_DOC_KEYS) & set(ALL_TRANSCRIPT_DOC_KEYS)
    assert shared == set(ALL_DOC_KEYS), "every solution document should have a pre-build counterpart"
    for key in shared:
        assert key in DOC_RENDERERS and key in TRANSCRIPT_DOC_RENDERERS


@pytest.mark.parametrize("key", DESIGN_KEYS)
def test_every_design_document_states_it_is_pre_build(key, snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], [key], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / f"{key}.md").read_text(encoding="utf-8")
    assert "pre-build version of this document" in text
    assert "exported solution zip" in text  # how to get the as-built version


def test_design_hld_scope_and_capabilities(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["hld"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "hld.md").read_text(encoding="utf-8")
    assert "# High-Level Design — Discovery Workshop 1 of 2 - Service Operations" in text
    assert "In scope — capability areas raised" in text
    assert "Job Scheduling" in text
    assert "Out of scope and deferred" in text
    assert "Open decisions" in text and "P1" in text


def test_design_data_dictionary_is_a_word_count_not_a_schema(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["data-dictionary"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "data-dictionary.md").read_text(encoding="utf-8")
    assert "Business vocabulary" in text
    assert "frequency count of what was said, not a proposed schema" in text
    assert "[Table / Column / Neither]" in text  # the decision is left to a human


def test_design_integration_lists_named_systems(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["integration"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "integration.md").read_text(encoding="utf-8")
    assert "Candidate interface inventory" in text
    assert "billing system" in text  # named in the workshop fixture


def test_design_security_surfaces_populations_and_leaves_the_matrix_empty(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["security"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "security.md").read_text(encoding="utf-8")
    assert "Role and privilege matrix" in text
    assert "Operations Director" in text  # candidate role from an attendee job title
    assert "[Consultant]" in text


def test_discovery_hygiene_reports_gaps_not_faults(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["hygiene"], ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    text = (tmp_path / "hygiene.md").read_text(encoding="utf-8")
    assert "# Discovery Hygiene Report" in text
    assert "DISC-301" in text  # A3 in the fixture has no due date
    assert "DISC-303" in text  # parked items are open decisions
    assert "Source record" in text


def test_design_documents_never_claim_a_built_system(snapshots, tmp_path: Path):
    """The whole contract: a pre-build document may *ask* for schema, never assert it.

    The signature of as-built metadata is a publisher-prefixed logical name
    (`abc_project`). Nothing in a discovery transcript can produce one, so its
    presence would mean a renderer had invented a component.
    """
    import re

    logical_name = re.compile(r"\b[a-z]{2,8}_[a-z][a-z0-9_]{2,}\b")
    render_transcript_documents(snapshots["workshop.txt"], list(DESIGN_KEYS), ["md"], tmp_path,
                                DocgenConfig(), no_llm=True)
    for key in DESIGN_KEYS:
        text = (tmp_path / f"{key}.md").read_text(encoding="utf-8")
        invented = logical_name.findall(text)
        assert not invented, f"{key}.md contains component-shaped names: {invented[:5]}"
        # and every design document defers its decisions rather than making them
        assert "[Consultant" in text, f"{key}.md makes design decisions on the consultant's behalf"


def test_docx_contains_title_and_tables(snapshots, tmp_path: Path):
    render_transcript_documents(snapshots["workshop.txt"], ["meeting-notes"], ["docx"], tmp_path,
                                DocgenConfig(), no_llm=True)
    docx = DocxDocument(str(tmp_path / "meeting-notes.docx"))
    all_text = "\n".join(p.text for p in docx.paragraphs)
    assert "Meeting Notes — Discovery Workshop 1 of 2 - Service Operations" in all_text
    assert len(docx.tables) >= 3
    cells = [c.text for t in docx.tables for row in t.rows for c in row.cells]
    assert "Rowan Ellisdale" in cells


def test_transcript_without_dialogue_still_renders(tmp_path: Path):
    """Every document, including the whole design set, survives an empty session."""
    cues = load_rules(DocgenConfig()).get("transcript_cues", {})
    source = tmp_path / "silent.txt"
    source.write_text("Just a note with no speakers at all.\n", encoding="utf-8")
    snapshot = parse_transcript(source, cues)
    written = render_transcript_documents(snapshot, list(TRANSCRIPT_DOC_RENDERERS), ["md"],
                                          tmp_path, DocgenConfig(), no_llm=True)
    assert len(written) == len(TRANSCRIPT_DOC_RENDERERS)
    requirements = (tmp_path / "requirements.md").read_text(encoding="utf-8")
    assert "No requirement candidates were detected" in requirements
    # a discovery that established nothing is reported as such, not as a clean bill
    hygiene = (tmp_path / "hygiene.md").read_text(encoding="utf-8")
    assert "Nothing was said about this" in hygiene


def test_run_all_transcript_end_to_end_offline(tmp_path: Path, monkeypatch):
    """`docgen all <transcript>` needs no network at all."""
    import socket

    def blocked(*args, **kwargs):  # pragma: no cover - only fires on regression
        raise AssertionError("offline transcript run attempted a network connection")

    monkeypatch.setattr(socket, "socket", blocked)
    monkeypatch.setattr(socket, "create_connection", blocked)

    run_all_transcript(TRANSCRIPTS / "workshop.txt", list(TRANSCRIPT_DOC_RENDERERS), ["md"],
                       tmp_path, DocgenConfig(), no_llm=True)
    assert (tmp_path / "transcript-snapshot.json").is_file()
    assert (tmp_path / "parse-warnings.md").is_file()
    for key in TRANSCRIPT_DOC_RENDERERS:
        assert (tmp_path / f"{key}.md").is_file()
